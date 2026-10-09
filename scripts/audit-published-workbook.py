"""Read every published service and compare its visible text with the original XLSX."""
import concurrent.futures
import datetime
import hashlib
import json
import sys
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook
from workbook_presentation_audit import verify_where_when, COLUMN

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1])
OUT = ROOT / "tmp/workbook-audit/published-final"
OUT.mkdir(parents=True, exist_ok=True)
sheet = load_workbook(SOURCE, data_only=True)["Catálogo consolidado"]
headers = {c.column: str(c.value).strip() for c in sheet[1] if c.value is not None}
rows = {str(row[0].value): {h: "" if row[c - 1].value is None else str(row[c - 1].value)
    for c, h in headers.items()} for row in sheet.iter_rows(min_row=2) if row[0].value is not None}
catalog = json.loads((ROOT / "tmp/workbook-audit/catalog-after.json").read_text(encoding="utf8"))

class Element:
    def __init__(self, tag="root", attrs=None):
        self.tag = tag
        self.attrs = dict(attrs or [])
        self.children = []
    def text(self):
        if "data-workbook-decoration" in self.attrs:
            return ""
        return "".join(c if isinstance(c, str) else c.text() for c in self.children)
    def all(self):
        yield self
        for c in self.children:
            if isinstance(c, Element): yield from c.all()

class Page(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Element()
        self.stack = [self.root]
    def handle_starttag(self, tag, attrs):
        node = Element(tag, attrs)
        self.stack[-1].children.append(node)
        if tag not in {"img", "br", "hr", "input", "meta", "link", "source", "wbr", "area", "embed", "param", "track", "col", "base"}:
            self.stack.append(node)
    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                break
    def handle_data(self, value): self.stack[-1].children.append(value)

def audit(service):
    url = "https://maisdigital.amargosa.ba.gov.br/servicos/" + service["slug"]
    result = {"id": service["id"], "sourceRow": service.get("sourceRow"), "url": url,
        "hasMirror": bool(service.get("workbookCells")), "differences": [], "checks": []}
    try:
        request = urllib.request.Request(url, headers={"User-Agent": "WorkbookContentAudit/2.0", "Cache-Control": "no-cache"})
        with urllib.request.urlopen(request, timeout=35) as response:
            html = response.read().decode("utf8")
            result["httpStatus"] = response.status
        (OUT / (service["slug"] + ".html")).write_text(html, encoding="utf8")
        p = Page(); p.feed(html)
        nodes = list(p.root.all())
        result["markedFields"] = sum("data-workbook-column" in n.attrs for n in nodes)
        if not result["hasMirror"]: return result
        cells = rows[str(service["sourceRow"])]
        def compare(column, actual):
            equal = bool(actual) and all(value == cells[column] for value in actual)
            result["checks"].append({"column": column, "expected": cells[column], "actual": actual, "equal": equal})
            if not equal: result["differences"].append(column)
        def section(id):
            return next((n for n in nodes if n.attrs.get("id") == id), Element())
        fields = ["Nome do serviço consolidado", "O que é", "Quem pode solicitar2", "Como solicitar - etapas",
            "Onde e quando solicitar2", "Custo", "Prazo", "Setor/Unidade responsável no SEI", "Telefone",
            "E-mail do setor/unidade", "Plataforma", "Link de acesso"]
        docs = [c for c in ["Documentos obrigatórios", "Documentos opcionais"] if cells[c]]
        if not docs and cells["Documentação necessária"]: docs = ["Documentação necessária"]
        fields += docs + [c for c in ["Secretaria responsável", "Ramal"] if cells[c]]
        for column in fields:
            presentations = [n.attrs["data-workbook-where-when-parts"] for n in nodes if "data-workbook-where-when-parts" in n.attrs]
            if column == COLUMN and presentations:
                try:
                    assert len(presentations) == 1
                    hour_nodes = {descendant for parent in nodes if parent.attrs.get("id") == "canais" or "service-request-presencial" in parent.attrs.get("class", "").split() for descendant in parent.all()}
                    slices = [(n.attrs.get("data-workbook-start"), n.attrs.get("data-workbook-end"), n.text(), "hours" if n in hour_nodes else "intro") for n in nodes if n.attrs.get("data-workbook-column") == COLUMN]
                    visible_fields = {c: [n.text() for n in nodes if n.attrs.get("data-workbook-column") == c] for c in fields}
                    verify_where_when(cells, presentations[0], slices, visible_fields)
                    result["checks"].append({"column": column, "equal": True, "presentation": "source slices and equivalent contact fields verified"})
                except AssertionError as error:
                    result["differences"].append(column)
                    result["checks"].append({"column": column, "equal": False, "error": str(error)})
                continue
            marked = [n.text() for n in nodes if n.attrs.get("data-workbook-column") == column]
            if marked:
                compare(column, marked); continue
            if column == "Nome do serviço consolidado":
                actual = [n.text() for n in nodes if n.tag == "h1"]
            elif column in {"O que é", "Quem pode solicitar2", "Onde e quando solicitar2"}:
                id = {"O que é": "o-que-e", "Quem pode solicitar2": "quem-pode", "Onde e quando solicitar2": "onde-quando"}[column]
                actual = [n.text() for n in section(id).all() if n.tag == "p"][:1]
            elif column == "Como solicitar - etapas":
                actual = [n.text() for n in section("como-solicitar").all() if n.tag == "ol"]
            elif column in docs:
                items = [n.text() for n in section("documentos").all() if n.tag == "li"]
                actual = [items[docs.index(column)]] if len(items) > docs.index(column) else []
            elif column in {"Custo", "Prazo"}:
                items = [n.text() for n in section("informacoes").all() if n.tag == "strong"]
                index = 0 if column == "Custo" else 1
                actual = [items[index]] if len(items) > index else []
            else:
                # Old deployments lack source markers: look only in the existing contact elements.
                candidates = [n.text() for id in ["canais", "onde-quando"] for n in section(id).all()
                    if n.tag in {"strong", "small", "a", "dd", "span"}]
                actual = [cells[column]] if cells[column] in candidates else candidates
            compare(column, actual)
        article = section("conteudo-servico")
        header = next((n for n in article.all() if n.tag == "header"), Element())
        compare("Público", [n.text() for n in header.all() if n.tag == "small"][:1])
        result["extraLegislation"] = bool(section("legislacao").children)
    except Exception as error:
        result["error"] = str(error)
    return result

with concurrent.futures.ThreadPoolExecutor(max_workers=6) as pool:
    results = list(pool.map(audit, catalog))
summary = {"checkedAt": datetime.datetime.now(datetime.timezone.utc).isoformat(),
    "sourceSha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(), "sourceRows": len(rows),
    "pages": len(results), "matchedPages": sum(r["hasMirror"] for r in results),
    "exceptions": sum(not r["hasMirror"] for r in results),
    "errors": sum("error" in r for r in results),
    "pagesWithDifferences": sum(bool(r["differences"]) for r in results),
    "fieldDifferences": sum(len(r["differences"]) for r in results),
    "pagesWithSourceMarkers": sum(bool(r.get("markedFields")) for r in results)}
(OUT / "comparison.json").write_text(json.dumps({"summary": summary, "pages": results}, ensure_ascii=False, indent=2), encoding="utf8")
print(json.dumps(summary))
