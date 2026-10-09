"""Verify visible text in the existing page elements against the original XLSX."""
import json
import sys
from html.parser import HTMLParser
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
folder = root / "tmp/workbook-audit"
source = Path(sys.argv[1]) if len(sys.argv) > 1 else root / "tmp/catalog-current.xlsx"
sheet = load_workbook(source, data_only=True)["Catálogo consolidado"]
headers = {c.column: str(c.value).strip() for c in sheet[1] if c.value is not None}
original = {str(row[0].value): {header: "" if row[col - 1].value is None else str(row[col - 1].value)
    for col, header in headers.items()} for row in sheet.iter_rows(min_row=2) if row[0].value is not None}
copied = json.loads((root / "app/workbook-service-data.json").read_text(encoding="utf8"))
assert set(copied["rows"]) == set(original)
for row_id, row in copied["rows"].items():
    assert row["cells"] == original[row_id], f"Source cells changed: {row_id}"

class Page(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack = []; self.fields = {}; self.ids = []; self.classes = []; self.headings = []
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        self.classes.extend(attrs.get("class", "").split())
        if attrs.get("id"): self.ids.append(attrs["id"])
        field = attrs.get("data-workbook-column")
        value = [] if field else None
        if field: self.fields.setdefault(field, []).append(value)
        if tag not in {"img", "br", "hr", "input", "meta", "link", "source", "wbr", "area"}:
            self.stack.append((tag, value, "data-workbook-decoration" in attrs))
    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, -1, -1):
            if self.stack[i][0] == tag:
                del self.stack[i:]; break
    def handle_data(self, value):
        if any(frame[2] for frame in self.stack): return
        for _, field, _ in self.stack:
            if field is not None: field.append(value)

catalog = {s["id"]: s for s in json.loads((folder / "catalog-after.json").read_text(encoding="utf8"))}
pages = json.loads((folder / "rendered-pages.json").read_text(encoding="utf8"))
count = 0
for service_id, html in pages.items():
    page = Page(); page.feed(html)
    cells = original[str(catalog[service_id]["sourceRow"])]
    expected = ["Nome do serviço consolidado", "O que é", "Quem pode solicitar2", "Como solicitar - etapas",
        "Onde e quando solicitar2", "Custo", "Prazo", "Setor/Unidade responsável no SEI", "Telefone",
        "E-mail do setor/unidade", "Plataforma", "Link de acesso"]
    documents = [c for c in ["Documentos obrigatórios", "Documentos opcionais"] if cells[c]]
    if not documents and cells["Documentação necessária"]: documents = ["Documentação necessária"]
    expected += documents
    expected += [c for c in ["Secretaria responsável", "Ramal"] if cells[c]]
    assert set(page.fields) == set(expected), f"Missing/extra displayed fields: {service_id}: {set(page.fields) ^ set(expected)}"
    for column in expected:
        for value in page.fields[column]:
            assert "".join(value) == cells[column], f"Changed visible text: {service_id}: {column}: {''.join(value)!r}"
            count += 1
    assert "legislacao" not in page.ids, f"Extra legislation: {service_id}"
    for css in ["service-steps"]:
        assert css in page.classes, f"Original component missing: {service_id}: {css}"
    assert "service-channel-list" in page.classes or "service-request-presencial" in page.classes, f"Contact card missing: {service_id}"
    assert '<section id="documentos"><h2>' in html
    assert '<section id="documentos"><h2>Documentos necessários</h2><ul>' in html or service_id == "planilha-27-comunicacao-de-animais-soltos"
    assert "Documentação necessária</h3>" not in html, f"Duplicated document section: {service_id}"
result = {"sourceRows": len(original), "fullPages": len(pages), "visibleFieldOccurrences": count,
    "differences": 0, "originalComponents": "preserved"}
(folder / "render-verification.json").write_text(json.dumps(result), encoding="utf8")
print(json.dumps(result))
